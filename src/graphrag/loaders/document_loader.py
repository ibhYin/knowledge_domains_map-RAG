import logging
from pathlib import Path

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

SUPPORTED_EXTS = {".pdf", ".md", ".txt", ".docx"}


def _load_pdf(path: Path, doc_id: str) -> list[Document]:
    import pypdf

    docs = []
    with path.open("rb") as f:
        reader = pypdf.PdfReader(f)
        for page_no, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            docs.append(
                Document(
                    page_content=text,
                    metadata={"doc_id": doc_id, "source": path.name, "page": page_no},
                )
            )
    return docs


def _load_docx(path: Path, doc_id: str) -> list[Document]:
    import docx

    parts = []
    doc = docx.Document(str(path))
    body = doc.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                row_text = "，".join(c for c in cells if c)
                if row_text:
                    parts.append("表格行：" + row_text)
    return [
        Document(
            page_content="\n".join(parts),
            metadata={"doc_id": doc_id, "source": path.name},
        )
    ]


def _load_text(path: Path, doc_id: str) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".md":
        # 去掉 YAML frontmatter
        if text.startswith("---"):
            end = text.find("\n---", 3)
            if end != -1:
                text = text[end + 4 :].lstrip("\n")
    return [
        Document(
            page_content=text,
            metadata={"doc_id": doc_id, "source": path.name},
        )
    ]


def load_document(path: str | Path, doc_id: str) -> list[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文档格式: {ext}（支持 {sorted(SUPPORTED_EXTS)}）")
    if ext == ".pdf":
        return _load_pdf(path, doc_id)
    if ext == ".docx":
        return _load_docx(path, doc_id)
    return _load_text(path, doc_id)
